import re
from pathlib import Path

from pypdf import PdfReader


def clean_text(text: str) -> str:
    """Remove HTML tags, excessive whitespace, and other noise."""
    text = re.sub(r'<script.*?</script>', '', text, flags=re.DOTALL)
    text = re.sub(r'<[^>]+>', ' ', text)
    text = re.sub(r'https?://\S+', '', text)
    text = re.sub(r'\{[^}]{50,}\}', '', text)
    text = re.sub(r'a:\d+:\{.*?\}', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()


def extract_text_from_pdf(file_path: str) -> list[dict]:
    """Extract text from a PDF file."""
    reader = PdfReader(file_path)
    pages = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            pages.append({
                "page_number": i + 1,
                "text": text.strip(),
            })

    return pages


def extract_text_from_txt(file_path: str) -> list[dict]:
    """Extract text from a plain text or markdown file."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()

    if not text.strip():
        return []

    return [{"page_number": 1, "text": text.strip()}]


def extract_text_from_docx(file_path: str) -> list[dict]:
    """Extract text from a Word document."""
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)

    if not text.strip():
        return []

    return [{"page_number": 1, "text": text.strip()}]


def extract_text_from_excel(file_path: str) -> list[dict]:
    """Extract text from an Excel spreadsheet, preserving row structure with headers."""
    import openpyxl
    import csv

    ext = Path(file_path).suffix.lower()
    pages = []

    if ext == ".csv":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            all_rows = list(reader)

        if not all_rows:
            return []

        headers = all_rows[0]
        header_line = " | ".join(headers)
        data_rows = all_rows[1:]

        rows_per_chunk = 20
        for i in range(0, len(data_rows), rows_per_chunk):
            chunk_rows = data_rows[i:i + rows_per_chunk]
            formatted_rows = []
            for row in chunk_rows:
                row_parts = []
                for j, cell in enumerate(row):
                    if cell.strip():
                        col_name = headers[j] if j < len(headers) else f"Column {j+1}"
                        cleaned = clean_text(cell)
                        if cleaned and len(cleaned) < 500:
                            row_parts.append(f"{col_name}: {cleaned}")
                if row_parts:
                    formatted_rows.append(" | ".join(row_parts))

            if formatted_rows:
                text = f"Columns: {header_line}\n\n" + "\n".join(formatted_rows)
                pages.append({"page_number": 1, "text": text})

    else:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        for sheet_num, sheet in enumerate(wb.worksheets, 1):
            all_rows = []
            for row in sheet.iter_rows(values_only=True):
                cell_values = [str(cell) if cell is not None else "" for cell in row]
                if any(v.strip() for v in cell_values):
                    all_rows.append(cell_values)

            if not all_rows:
                continue

            headers = all_rows[0]
            header_line = " | ".join(headers)
            data_rows = all_rows[1:]

            rows_per_chunk = 20
            for i in range(0, len(data_rows), rows_per_chunk):
                chunk_rows = data_rows[i:i + rows_per_chunk]
                formatted_rows = []
                for row in chunk_rows:
                    row_parts = []
                    for j, cell in enumerate(row):
                        if cell.strip():
                            col_name = headers[j] if j < len(headers) else f"Column {j+1}"
                            cleaned = clean_text(cell)
                            if cleaned and len(cleaned) < 500:
                                row_parts.append(f"{col_name}: {cleaned}")
                    if row_parts:
                        formatted_rows.append(" | ".join(row_parts))

                if formatted_rows:
                    text = f"Sheet: {sheet.title}\nColumns: {header_line}\n\n" + "\n".join(formatted_rows)
                    pages.append({"page_number": sheet_num, "text": text})

        wb.close()

    return pages


def extract_text(file_path: str) -> list[dict]:
    """Extract text from a file based on its extension."""
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in (".txt", ".md"):
        return extract_text_from_txt(file_path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_path)
    elif ext in (".xlsx", ".xls", ".csv"):
        return extract_text_from_excel(file_path)
    else:
        return []